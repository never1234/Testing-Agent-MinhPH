"""One interface over Excel (.xlsx via openpyxl) and Google Sheets (gspread).

Book methods: sheet_names(), values(name), write(name, [(row, col, value)]), add_sheet(name, header),
backup(dest_dir), save(). Rows/cols are 1-based. Writes are buffered until save().
"""
from __future__ import annotations

import shutil
from datetime import date, datetime
from pathlib import Path

from .config import die, env, kit_path, warn


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


class ExcelBook:
    kind = "excel"

    def __init__(self, path: str):
        import openpyxl

        self.path = kit_path(path)
        if not self.path.exists():
            die(f"Excel file not found: {self.path}")
        lock = self.path.with_name("~$" + self.path.name)
        if lock.exists():
            die(f"{self.path.name} is open in Excel (found {lock.name}). Close it and retry.")
        self.wb = openpyxl.load_workbook(self.path)
        self._dirty = False

    def sheet_names(self) -> list[str]:
        return list(self.wb.sheetnames)

    def values(self, name: str) -> list[list[str]]:
        ws = self.wb[name]
        return [[_cell_str(c) for c in row] for row in ws.iter_rows(values_only=True)]

    def write(self, name: str, updates: list[tuple[int, int, object]]) -> None:
        ws = self.wb[name]
        for r, c, v in updates:
            ws.cell(row=r, column=c, value=v)
        if updates:
            self._dirty = True

    def add_sheet(self, name: str, header: list[str]) -> None:
        ws = self.wb.create_sheet(title=name)
        ws.append(header)
        self._dirty = True

    def backup(self, dest_dir: Path) -> Path:
        dest_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        dest = dest_dir / f"{self.path.stem}.bak-{stamp}{self.path.suffix}"
        shutil.copy2(self.path, dest)
        return dest

    def save(self) -> None:
        if self._dirty:
            self.wb.save(self.path)
            self._dirty = False


class GSheetBook:
    kind = "gsheet"

    def __init__(self, sheet_id: str):
        import gspread

        cred = kit_path(env("GOOGLE_SERVICE_ACCOUNT_JSON"))
        if not cred.exists():
            die(f"Service account JSON not found: {cred}")
        gc = gspread.service_account(filename=str(cred))
        try:
            self.sh = gc.open_by_key(sheet_id)
        except gspread.exceptions.APIError as e:
            die(f"Cannot open Google Sheet {sheet_id}: {e}. Did you share it with the service account email (Editor)?")
        self._ws = {ws.title: ws for ws in self.sh.worksheets()}
        self._pending: dict[str, list] = {}

    def sheet_names(self) -> list[str]:
        return list(self._ws)

    def values(self, name: str) -> list[list[str]]:
        return [[_cell_str(c) for c in row] for row in self._ws[name].get_all_values()]

    def write(self, name: str, updates: list[tuple[int, int, object]]) -> None:
        self._pending.setdefault(name, []).extend(
            (r, c, _cell_str(v) if isinstance(v, (date, datetime)) else v) for r, c, v in updates)

    def add_sheet(self, name: str, header: list[str]) -> None:
        ws = self.sh.add_worksheet(title=name, rows=500, cols=max(len(header), 10))
        self._ws[name] = ws
        self.write(name, [(1, i + 1, h) for i, h in enumerate(header)])

    def backup(self, dest_dir: Path) -> None:
        print("[INFO] Google Sheets keeps version history (File > Version history) - no local backup needed.")
        return None

    def save(self) -> None:
        from gspread.utils import ValueInputOption, rowcol_to_a1

        for name, ups in self._pending.items():
            if not ups:
                continue
            ws = self._ws[name]
            max_r = max(r for r, _, _ in ups)
            max_c = max(c for _, c, _ in ups)
            if max_r > ws.row_count:
                ws.add_rows(max_r - ws.row_count + 50)
            if max_c > ws.col_count:
                ws.add_cols(max_c - ws.col_count)
            data = [{"range": rowcol_to_a1(r, c), "values": [[v]]} for r, c, v in ups]
            for i in range(0, len(data), 500):
                ws.batch_update(data[i:i + 500], value_input_option=ValueInputOption.user_entered)
        self._pending = {}


def open_book(source: str, excel_path: str | None = None, gsheet_id: str | None = None):
    source = (source or "").lower()
    if source == "excel":
        if not excel_path:
            die("excel_path is not set in profile.yaml")
        return ExcelBook(excel_path)
    if source == "gsheet":
        if not gsheet_id:
            die("gsheet_id is not set in profile.yaml")
        return GSheetBook(gsheet_id)
    die(f"Unknown source '{source}' (use excel or gsheet)")


def open_testcase_book(prof: dict):
    tc = prof["testcases"]
    return open_book(tc.get("source", "excel"), tc.get("excel_path"), tc.get("gsheet_id"))


# ---------- table helpers ----------

def locate_table(values: list[list[str]], columns: dict, scan_rows: int = 40):
    """Find the header row (the row containing the ID column name).
    Returns (header_row_index_0based, {key: col_index_0based or None}) or (None, None)."""
    id_name = str(columns.get("id", "ID")).strip().lower()
    for r, row in enumerate(values[:scan_rows]):
        norm = [str(c).strip().lower() for c in row]
        if id_name in norm:
            idx = {}
            for key, name in columns.items():
                n = str(name or "").strip().lower()
                idx[key] = norm.index(n) if n and n in norm else None
            return r, idx
    return None, None


def iter_table(values: list[list[str]], header_r: int, idx: dict):
    """Yield dicts {'_row': 1-based row number, key: value} for rows with a non-empty ID."""
    id_c = idx["id"]
    for r in range(header_r + 1, len(values)):
        row = values[r]
        rid = row[id_c].strip() if id_c < len(row) else ""
        if not rid:
            continue
        rec = {"_row": r + 1}
        for key, c in idx.items():
            if c is not None:
                rec[key] = row[c].strip() if c < len(row) else ""
        yield rec


def build_id_index(book, prof: dict) -> dict:
    """Map upper(ID) -> (sheet_name, record, idx) across all non-skipped sheets."""
    cols = prof["testcases"]["columns"]
    skip = {s.lower() for s in prof["testcases"].get("skip_sheets", [])}
    index = {}
    for name in book.sheet_names():
        if name.lower() in skip:
            continue
        values = book.values(name)
        hr, idx = locate_table(values, cols)
        if hr is None:
            continue
        for rec in iter_table(values, hr, idx):
            key = rec["id"].upper()
            if key in index:
                warn(f"Duplicate ID {rec['id']} in sheets '{index[key][0]}' and '{name}' - using the first")
                continue
            index[key] = (name, rec, idx)
    return index
